#!/usr/bin/env python3
"""Sync groundwater source files and generated indexes with Google Drive.

This script is intended for GitHub Actions. It reads source files from Google
Drive, generates stable indexes, writes index artifacts back to Drive, and
leaves local copies for the public-site build step.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import mimetypes
import re
import sys
import unicodedata
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

DRIVE_FIELDS = (
    "files(id,name,mimeType,modifiedTime,size,md5Checksum,parents,webViewLink),"
    "nextPageToken"
)

INDEX_FILENAMES = {
    "well": "well-index.json",
    "station": "station-index.json",
    "wellWarnings": "warnings.csv",
    "attachments": "water-right-attachment-index.json",
    "pumping": "pumping-index.json",
    "pumpingMonth": "pumping-month-index.json",
    "pumpingWarnings": "pumping-warnings.json",
    "sync": "sync-index.json",
    "summary": "sync-summary.md",
}

EXCEL_MIME_TYPES = {
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "application/vnd.ms-excel",
    "application/vnd.google-apps.spreadsheet",
}


def normalize_text(value: object) -> str:
    if value is None:
        return ""
    text = unicodedata.normalize("NFKC", str(value))
    return re.sub(r"\s+", " ", text).strip()


def normalize_key(value: object) -> str:
    return normalize_text(value).upper().replace(" ", "")


def extract_water_right_key(value: object) -> str:
    text = normalize_key(value)
    match = re.search(r"([A-Z])(\d{6,7})", text)
    if not match:
        return ""
    return f"{match.group(1)}{match.group(2).zfill(7)}"


def record_hash(record: dict[str, Any]) -> str:
    payload = json.dumps(record, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def parse_source_date(name: str) -> str:
    """Return a comparable YYYY-MM-DD string when a date can be inferred."""
    text = unicodedata.normalize("NFKC", name)
    patterns = [
        r"(?P<roc>\d{3})[年\-_./]?(?P<m>\d{1,2})[月\-_./]?(?P<d>\d{1,2})日?",
        r"(?P<y>20\d{2})[年\-_./](?P<m>\d{1,2})[月\-_./](?P<d>\d{1,2})日?",
        r"(?P<roc>\d{3})年度?",
    ]
    for pattern in patterns:
        match = re.search(pattern, text)
        if not match:
            continue
        groups = match.groupdict()
        if groups.get("roc"):
            year = int(groups["roc"]) + 1911
        else:
            year = int(groups["y"])
        month = int(groups.get("m") or 1)
        day = int(groups.get("d") or 1)
        try:
            return datetime(year, month, day, tzinfo=timezone.utc).date().isoformat()
        except ValueError:
            return ""
    return ""


def is_excel_file(file: dict[str, Any]) -> bool:
    name = file.get("name", "")
    mime_type = file.get("mimeType", "")
    return (
        mime_type in EXCEL_MIME_TYPES
        or name.lower().endswith((".xlsx", ".xlsm", ".xls"))
    )


def is_registry_candidate(file: dict[str, Any]) -> bool:
    name = file.get("name", "")
    return is_excel_file(file) and ("一覽表" in name or "抽水井" in name or "井籍" in name)


def is_pumping_candidate(file: dict[str, Any]) -> bool:
    name = file.get("name", "")
    return (
        "抽水" in name
        and ("紀錄" in name or "history" in name.lower() or "pumping" in name.lower())
        and (is_excel_file(file) or name.lower().endswith(".json"))
    )


def drive_service(service_account_path: Path):
    from google.oauth2 import service_account
    from googleapiclient.discovery import build

    credentials = service_account.Credentials.from_service_account_file(
        service_account_path,
        scopes=["https://www.googleapis.com/auth/drive"],
    )
    return build("drive", "v3", credentials=credentials, cache_discovery=False)


def list_children(service, folder_id: str) -> list[dict[str, Any]]:
    files: list[dict[str, Any]] = []
    page_token = None
    while True:
        response = service.files().list(
            q=f"'{folder_id}' in parents and trashed=false",
            fields=DRIVE_FIELDS,
            pageToken=page_token,
            pageSize=1000,
            supportsAllDrives=True,
            includeItemsFromAllDrives=True,
        ).execute()
        files.extend(response.get("files", []))
        page_token = response.get("nextPageToken")
        if not page_token:
            return files


def list_tree(service, folder_id: str) -> list[dict[str, Any]]:
    seen_folders: set[str] = set()
    found: list[dict[str, Any]] = []

    def visit(current_folder_id: str) -> None:
        if current_folder_id in seen_folders:
            return
        seen_folders.add(current_folder_id)
        for item in list_children(service, current_folder_id):
            found.append(item)
            if item.get("mimeType") == "application/vnd.google-apps.folder":
                visit(item["id"])

    visit(folder_id)
    return found


def newest_file(files: list[dict[str, Any]], predicate) -> dict[str, Any] | None:
    candidates = [file for file in files if predicate(file)]
    if not candidates:
        return None

    def sort_key(file: dict[str, Any]):
        source_date = parse_source_date(file.get("name", ""))
        modified = file.get("modifiedTime") or ""
        size = int(file.get("size") or 0)
        return (source_date or "0000-00-00", modified, size, file.get("name", ""))

    return sorted(candidates, key=sort_key, reverse=True)[0]


def file_metadata_state(file: dict[str, Any]) -> dict[str, Any]:
    return {
        "id": file.get("id"),
        "name": file.get("name"),
        "mimeType": file.get("mimeType"),
        "modifiedTime": file.get("modifiedTime"),
        "size": file.get("size", ""),
        "md5Checksum": file.get("md5Checksum", ""),
        "sourceDate": parse_source_date(file.get("name", "")),
        "webViewLink": file.get("webViewLink", ""),
    }


def find_child_by_name(service, folder_id: str, name: str) -> dict[str, Any] | None:
    escaped = name.replace("'", "\\'")
    page_token = None
    while True:
        response = service.files().list(
            q=f"'{folder_id}' in parents and name='{escaped}' and trashed=false",
            fields=DRIVE_FIELDS,
            pageToken=page_token,
            pageSize=100,
            supportsAllDrives=True,
            includeItemsFromAllDrives=True,
        ).execute()
        files = response.get("files", [])
        if files:
            return sorted(files, key=lambda item: item.get("modifiedTime", ""), reverse=True)[0]
        page_token = response.get("nextPageToken")
        if not page_token:
            return None


def download_file(service, file: dict[str, Any], target: Path) -> Path:
    from googleapiclient.http import MediaIoBaseDownload

    target.parent.mkdir(parents=True, exist_ok=True)
    mime_type = file.get("mimeType", "")
    if mime_type == "application/vnd.google-apps.spreadsheet":
        request = service.files().export_media(
            fileId=file["id"],
            mimeType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
    else:
        request = service.files().get_media(fileId=file["id"], supportsAllDrives=True)

    with target.open("wb") as fh:
        downloader = MediaIoBaseDownload(fh, request)
        done = False
        while not done:
            _, done = downloader.next_chunk()
    return target


def download_index_if_exists(service, folder_id: str, name: str, target: Path) -> bool:
    file = find_child_by_name(service, folder_id, name)
    if not file:
        return False
    download_file(service, file, target)
    return True


def upload_or_update(service, folder_id: str, path: Path, mime_type: str | None = None) -> dict[str, Any]:
    from googleapiclient.http import MediaFileUpload

    mime_type = mime_type or mimetypes.guess_type(path.name)[0] or "application/octet-stream"
    media = MediaFileUpload(str(path), mimetype=mime_type, resumable=False)
    existing = find_child_by_name(service, folder_id, path.name)
    fields = "id,name,modifiedTime,size,md5Checksum,webViewLink"
    if existing:
        return service.files().update(
            fileId=existing["id"],
            media_body=media,
            fields=fields,
            supportsAllDrives=True,
        ).execute()
    return service.files().create(
        body={"name": path.name, "parents": [folder_id]},
        media_body=media,
        fields=fields,
        supportsAllDrives=True,
    ).execute()


def upload_indexes(service, folder_id: str, names: list[str], work_dir: Path, label: str, notes: list[str]) -> None:
    for name in names:
        path = work_dir / name
        if not path.exists():
            continue
        try:
            upload_or_update(service, folder_id, path)
        except Exception as exc:
            notes.append(f"Could not upload {label} index {name} to Drive folder {folder_id}: {exc}")
            print(f"warning: could not upload {label} index {name}: {exc}", file=sys.stderr)


def find_header_row(sheet, required: set[str], max_scan: int = 20) -> tuple[int, dict[str, int]]:
    for row_index in range(1, max_scan + 1):
        values = [normalize_text(cell.value) for cell in sheet[row_index]]
        positions = {name: i for i, name in enumerate(values) if name}
        if required.issubset(positions):
            return row_index, positions
    raise ValueError(f"Cannot find header row with required columns: {', '.join(sorted(required))}")


def unique_header_names(raw_names: list[str]) -> list[str]:
    names: list[str] = []
    counts: dict[str, int] = {}
    previous = ""
    fillable_period_headers = {"核准水權年限", "下次申請時間"}
    for raw_name in raw_names:
        name = raw_name
        if not name and previous in fillable_period_headers:
            name = previous
        if name:
            previous = name
            counts[name] = counts.get(name, 0) + 1
            if counts[name] > 1:
                name = f"{name}_{counts[previous]}"
        names.append(name)
    return names


def image_anchor_cell(image) -> tuple[int | None, int | None, str]:
    marker = getattr(getattr(image, "anchor", None), "_from", None)
    if marker is None:
        return None, None, ""
    row = getattr(marker, "row", None)
    col = getattr(marker, "col", None)
    if row is None or col is None:
        return None, None, ""
    row_number = int(row) + 1
    col_number = int(col) + 1
    return row_number, col_number, f"R{row_number}C{col_number}"


def image_bytes(image) -> bytes:
    data = image._data()
    return data if isinstance(data, bytes) else bytes(data)


def find_nearby_image_key(sheet, row_number: int, col_number: int) -> str:
    row_start = max(1, row_number - 14)
    row_end = min(row_number + 2, sheet.max_row or row_number + 2)
    col_start = max(1, col_number - 10)
    col_end = min(col_number + 8, sheet.max_column or col_number + 8)
    candidates: list[tuple[int, int, str]] = []
    for row in range(row_start, row_end + 1):
        for col in range(col_start, col_end + 1):
            key = extract_water_right_key(sheet.cell(row, col).value)
            if key:
                distance = abs(row_number - row) + abs(col_number - col)
                candidates.append((distance, row, key))
    if not candidates:
        return ""
    candidates.sort(key=lambda item: (item[0], -item[1], item[2]))
    return candidates[0][2]


def collect_embedded_images(workbook) -> tuple[dict[str, list[dict[str, str]]], list[dict[str, str]]]:
    by_key: dict[str, list[dict[str, str]]] = defaultdict(list)
    warnings: list[dict[str, str]] = []
    for sheet in workbook.worksheets:
        for image in getattr(sheet, "_images", []):
            row_number, col_number, anchor = image_anchor_cell(image)
            try:
                raw = image_bytes(image)
            except Exception as exc:
                warnings.append({
                    "wellKey": "",
                    "station": "",
                    "sourceRow": "",
                    "reason": f"embedded_photo_unreadable:{sheet.title}:{anchor or 'unknown'}:{exc}",
                })
                continue
            item = {
                "sheet": sheet.title,
                "anchor": anchor,
                "imageHash": hashlib.sha256(raw).hexdigest(),
                "extension": normalize_text(getattr(image, "format", "")),
                "size": str(len(raw)),
            }
            if row_number is None:
                warnings.append({
                    "wellKey": "",
                    "station": sheet.title,
                    "sourceRow": "",
                    "reason": "embedded_photo_without_anchor",
                })
                continue
            well_key = find_nearby_image_key(sheet, row_number, col_number or 1)
            if not well_key:
                warnings.append({
                    "wellKey": "",
                    "station": sheet.title,
                    "sourceRow": str(row_number),
                    "reason": f"embedded_photo_unmatched:{sheet.title}:{anchor}",
                })
                continue
            item["matchedWellKey"] = well_key
            by_key[well_key].append(item)
    return by_key, warnings


def combined_photo_hash(images: list[dict[str, str]]) -> str:
    if not images:
        return ""
    payload = json.dumps(
        sorted(images, key=lambda item: (item["sheet"], item["anchor"], item["imageHash"])),
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
    )
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def classify_excel(excel_path: Path) -> tuple[list[dict[str, Any]], dict[str, list[dict[str, Any]]], list[dict[str, str]]]:
    from openpyxl import load_workbook

    workbook = load_workbook(excel_path, data_only=True, read_only=False)
    if "總表" not in workbook.sheetnames:
        raise ValueError("Workbook must contain a 總表 sheet")

    embedded_images_by_key, image_warnings = collect_embedded_images(workbook)
    sheet = workbook["總表"]
    header_row, _ = find_header_row(sheet, {"站別", "水權狀號"})
    header_names = unique_header_names([normalize_text(cell.value) for cell in sheet[header_row]])
    records: list[dict[str, Any]] = []
    warnings: list[dict[str, str]] = []
    seen: dict[str, int] = {}

    for row_number, row in enumerate(sheet.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1):
        source = {
            header_names[i]: normalize_text(value)
            for i, value in enumerate(row)
            if i < len(header_names) and header_names[i]
        }
        if not any(source.values()):
            continue
        well_key = extract_water_right_key(source.get("水權狀號")) or normalize_key(source.get("水權狀號"))
        station = normalize_text(source.get("站別"))
        ambiguous = ""
        if not well_key:
            fallback = "|".join([station, source.get("井別", ""), source.get("井址", ""), str(row_number)])
            well_key = "ROW-" + hashlib.sha1(fallback.encode("utf-8")).hexdigest()[:12].upper()
            ambiguous = "missing_water_right_no"
        elif well_key in seen:
            ambiguous = "duplicate_water_right_no"

        embedded_images = embedded_images_by_key.get(well_key, [])
        source_hash = record_hash(source)
        photo_hash = combined_photo_hash(embedded_images)
        record = {
            "wellKey": well_key,
            "station": station,
            "sourceRow": str(row_number),
            "ambiguous": ambiguous,
            "rowHash": source_hash,
            "photoHash": photo_hash,
            "changeHash": record_hash({"sourceHash": source_hash, "photoHash": photo_hash}),
            "embeddedPhotos": embedded_images,
            "source": source,
        }
        records.append(record)
        seen[well_key] = seen.get(well_key, 0) + 1
        if ambiguous:
            warnings.append({
                "wellKey": well_key,
                "station": station,
                "sourceRow": str(row_number),
                "reason": ambiguous,
            })

    warnings.extend(image_warnings)
    stations: dict[str, list[dict[str, Any]]] = {}
    for record in records:
        stations.setdefault(record["station"] or "_未分站", []).append(record)
    return records, stations, warnings


def write_json(path: Path, value: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def read_json_if_exists(path: Path, default: Any) -> Any:
    if not path.exists():
        return default
    return json.loads(path.read_text(encoding="utf-8"))


def write_warnings_csv(path: Path, warnings: list[dict[str, str]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8-sig") as fh:
        writer = csv.DictWriter(fh, fieldnames=["wellKey", "station", "sourceRow", "reason"])
        writer.writeheader()
        writer.writerows(warnings)


def water_right_attachment_index(files: list[dict[str, Any]]) -> tuple[list[dict[str, Any]], list[dict[str, str]]]:
    grouped: dict[str, list[dict[str, Any]]] = defaultdict(list)
    warnings: list[dict[str, str]] = []
    for file in files:
        if file.get("mimeType") == "application/vnd.google-apps.folder":
            continue
        name = file.get("name", "")
        if not name.lower().endswith(".pdf") and "水權" not in name:
            continue
        key = extract_water_right_key(name)
        if not key:
            warnings.append({
                "wellKey": "",
                "station": "",
                "sourceRow": "",
                "reason": f"water_right_file_unmatched:{name}",
            })
            continue
        grouped[key].append(file)

    index: list[dict[str, Any]] = []
    for key, matches in grouped.items():
        matches.sort(key=lambda item: (item.get("modifiedTime", ""), int(item.get("size") or 0), item.get("name", "")), reverse=True)
        newest = matches[0]
        if len(matches) > 1 and matches[0].get("modifiedTime") == matches[1].get("modifiedTime"):
            warnings.append({
                "wellKey": key,
                "station": "",
                "sourceRow": "",
                "reason": f"water_right_file_tie:{matches[0].get('name')}:{matches[1].get('name')}",
            })
        index.append({
            "wellKey": key,
            "fileId": newest.get("id"),
            "name": newest.get("name"),
            "mimeType": newest.get("mimeType"),
            "modifiedTime": newest.get("modifiedTime"),
            "size": newest.get("size", ""),
            "md5Checksum": newest.get("md5Checksum", ""),
            "webViewLink": newest.get("webViewLink", ""),
            "candidateCount": len(matches),
        })
    index.sort(key=lambda item: item["wellKey"])
    return index, warnings


def build_pumping_indexes(source_path: Path, out_dir: Path) -> dict[str, int]:
    data = json.loads(source_path.read_text(encoding="utf-8"))
    records = data.get("records", [])
    yearly = []
    monthly = []
    warnings = []

    for source_index, record in enumerate(records):
        water_right_no = normalize_key(record.get("waterRightNo"))
        year = record.get("yearMinguo")
        months = record.get("monthlyM3") or []
        if not water_right_no:
            warnings.append({
                "sourceIndex": source_index,
                "waterRightNo": "",
                "yearMinguo": year,
                "reason": "missing_water_right_no",
            })
            continue
        if not isinstance(year, int):
            warnings.append({
                "sourceIndex": source_index,
                "waterRightNo": water_right_no,
                "yearMinguo": year,
                "reason": "invalid_year",
            })
            continue
        normalized = {
            "waterRightNo": water_right_no,
            "yearMinguo": year,
            "station": normalize_text(record.get("station")),
            "wellName": normalize_text(record.get("wellName")),
            "authority": normalize_text(record.get("authority")),
            "monthlyM3": months,
            "sourceTotalM3": record.get("sourceTotalM3"),
        }
        yearly.append({
            **normalized,
            "yearlyKey": f"{water_right_no}:{year}",
            "recordHash": record_hash(normalized),
            "sourceIndex": source_index,
        })
        for month_index in range(12):
            value = months[month_index] if month_index < len(months) else None
            monthly.append({
                "monthlyKey": f"{water_right_no}:{year}:{month_index + 1:02d}",
                "waterRightNo": water_right_no,
                "yearMinguo": year,
                "month": month_index + 1,
                "m3": value,
            })

    write_json(out_dir / INDEX_FILENAMES["pumping"], yearly)
    write_json(out_dir / INDEX_FILENAMES["pumpingMonth"], monthly)
    write_json(out_dir / INDEX_FILENAMES["pumpingWarnings"], warnings)
    return {"records": len(yearly), "monthlyRecords": len(monthly), "warnings": len(warnings)}


def should_parse_registry(previous_sync: dict[str, Any], current_file: dict[str, Any]) -> bool:
    previous = ((previous_sync.get("sources") or {}).get("registryExcel") or {})
    current = file_metadata_state(current_file)
    fields = ["id", "sourceDate", "modifiedTime", "size", "md5Checksum"]
    return any(str(previous.get(field, "")) != str(current.get(field, "")) for field in fields)


def write_summary(path: Path, summary: dict[str, Any]) -> None:
    lines = [
        "# Groundwater Drive Sync Summary",
        "",
        f"- Generated at: {summary['generatedAt']}",
        f"- Sync scope: {summary['syncScope']}",
        f"- Registry Excel parsed: {summary['registry']['parsed']}",
        f"- Well records: {summary['registry']['records']}",
        f"- Well warnings: {summary['registry']['warnings']}",
        f"- Water-right attachment records: {summary['waterRights']['records']}",
        f"- Water-right attachment warnings: {summary['waterRights']['warnings']}",
        f"- Pumping records: {summary['pumping']['records']}",
        f"- Pumping monthly records: {summary['pumping']['monthlyRecords']}",
        f"- Pumping warnings: {summary['pumping']['warnings']}",
        "",
    ]
    if summary.get("notes"):
        lines.append("## Notes")
        lines.append("")
        lines.extend(f"- {note}" for note in summary["notes"])
        lines.append("")
    path.write_text("\n".join(lines), encoding="utf-8")


def main() -> int:
    parser = argparse.ArgumentParser(description="Sync groundwater Google Drive indexes for GitHub Actions.")
    parser.add_argument("--service-account", required=True, type=Path)
    parser.add_argument("--sync-scope", default="all", choices=["all", "wells", "pumping"])
    parser.add_argument("--groundwater-root-folder-id", required=True)
    parser.add_argument("--registry-folder-id", default="")
    parser.add_argument("--well-index-folder-id", required=True)
    parser.add_argument("--pumping-index-folder-id", required=True)
    parser.add_argument("--water-right-folder-id", required=True)
    parser.add_argument("--work-dir", required=True, type=Path)
    args = parser.parse_args()

    args.work_dir.mkdir(parents=True, exist_ok=True)
    service = drive_service(args.service_account)
    notes: list[str] = []

    sync_index_path = args.work_dir / INDEX_FILENAMES["sync"]
    previous_sync_exists = download_index_if_exists(
        service,
        args.well_index_folder_id,
        INDEX_FILENAMES["sync"],
        sync_index_path,
    )
    previous_sync = read_json_if_exists(sync_index_path, {}) if previous_sync_exists else {}

    root_files = list_tree(service, args.groundwater_root_folder_id)
    registry_files = list_tree(service, args.registry_folder_id) if args.registry_folder_id else root_files
    registry_file = newest_file(registry_files, is_registry_candidate)
    if not registry_file and args.sync_scope in {"all", "wells"}:
        raise RuntimeError("No groundwater registry Excel candidate found in Drive")

    well_records: list[dict[str, Any]] = []
    station_records: dict[str, list[dict[str, Any]]] = {}
    well_warnings: list[dict[str, str]] = []
    registry_parsed = False

    if args.sync_scope in {"all", "wells"} and registry_file:
        if should_parse_registry(previous_sync, registry_file):
            registry_target = args.work_dir / "source-registry.xlsx"
            download_file(service, registry_file, registry_target)
            well_records, station_records, well_warnings = classify_excel(registry_target)
            write_json(args.work_dir / INDEX_FILENAMES["well"], well_records)
            write_json(args.work_dir / INDEX_FILENAMES["station"], station_records)
            write_warnings_csv(args.work_dir / INDEX_FILENAMES["wellWarnings"], well_warnings)
            registry_parsed = True
        else:
            notes.append("Registry Excel metadata unchanged; reused Drive-stored well indexes.")
            for name in [INDEX_FILENAMES["well"], INDEX_FILENAMES["station"], INDEX_FILENAMES["wellWarnings"]]:
                if not download_index_if_exists(service, args.well_index_folder_id, name, args.work_dir / name):
                    raise RuntimeError(f"Registry unchanged but missing Drive index file: {name}")
            well_records = read_json_if_exists(args.work_dir / INDEX_FILENAMES["well"], [])
            station_records = read_json_if_exists(args.work_dir / INDEX_FILENAMES["station"], {})
            with (args.work_dir / INDEX_FILENAMES["wellWarnings"]).open("r", encoding="utf-8-sig") as fh:
                well_warnings = list(csv.DictReader(fh))
    else:
        for name in [INDEX_FILENAMES["well"], INDEX_FILENAMES["station"], INDEX_FILENAMES["wellWarnings"]]:
            download_index_if_exists(service, args.well_index_folder_id, name, args.work_dir / name)
        well_records = read_json_if_exists(args.work_dir / INDEX_FILENAMES["well"], [])
        station_records = read_json_if_exists(args.work_dir / INDEX_FILENAMES["station"], {})

    water_right_files = list_tree(service, args.water_right_folder_id)
    attachment_records, attachment_warnings = water_right_attachment_index(water_right_files)
    write_json(args.work_dir / INDEX_FILENAMES["attachments"], attachment_records)
    if attachment_warnings:
        well_warnings.extend(attachment_warnings)
        write_warnings_csv(args.work_dir / INDEX_FILENAMES["wellWarnings"], well_warnings)

    pumping_stats = {"records": 0, "monthlyRecords": 0, "warnings": 0}
    if args.sync_scope in {"all", "pumping"}:
        pumping_file = newest_file(root_files, is_pumping_candidate)
        if pumping_file and pumping_file.get("name", "").lower().endswith(".json"):
            target = args.work_dir / "source-pumping-history.json"
            download_file(service, pumping_file, target)
            pumping_stats = build_pumping_indexes(target, args.work_dir)
        elif pumping_file:
            notes.append(
                "Pumping source appears to be an Excel file; add a pumping Excel parser before replacing public pumping data."
            )
        else:
            notes.append("No pumping source candidate found; reused Drive-stored pumping indexes when available.")
        for name in [INDEX_FILENAMES["pumping"], INDEX_FILENAMES["pumpingMonth"], INDEX_FILENAMES["pumpingWarnings"]]:
            path = args.work_dir / name
            if not path.exists():
                download_index_if_exists(service, args.pumping_index_folder_id, name, path)
        existing_pumping = read_json_if_exists(args.work_dir / INDEX_FILENAMES["pumping"], [])
        existing_monthly = read_json_if_exists(args.work_dir / INDEX_FILENAMES["pumpingMonth"], [])
        existing_warnings = read_json_if_exists(args.work_dir / INDEX_FILENAMES["pumpingWarnings"], [])
        pumping_stats = {
            "records": len(existing_pumping),
            "monthlyRecords": len(existing_monthly),
            "warnings": len(existing_warnings),
        }
    else:
        for name in [INDEX_FILENAMES["pumping"], INDEX_FILENAMES["pumpingMonth"], INDEX_FILENAMES["pumpingWarnings"]]:
            download_index_if_exists(service, args.pumping_index_folder_id, name, args.work_dir / name)

    generated_at = datetime.now(timezone.utc).isoformat()
    sync_index = {
        "schemaVersion": 1,
        "generatedAt": generated_at,
        "syncScope": args.sync_scope,
        "sources": {
            "registryExcel": file_metadata_state(registry_file) if registry_file else {},
            "registryFolderId": args.registry_folder_id or args.groundwater_root_folder_id,
            "waterRightFolderId": args.water_right_folder_id,
            "groundwaterRootFolderId": args.groundwater_root_folder_id,
        },
        "outputs": {
            "wellRecords": len(well_records),
            "stationGroups": len(station_records),
            "wellWarnings": len(well_warnings),
            "waterRightAttachmentRecords": len(attachment_records),
            "waterRightAttachmentWarnings": len(attachment_warnings),
            "pumpingRecords": pumping_stats["records"],
            "pumpingMonthlyRecords": pumping_stats["monthlyRecords"],
            "pumpingWarnings": pumping_stats["warnings"],
        },
        "notes": notes,
    }
    write_json(args.work_dir / INDEX_FILENAMES["sync"], sync_index)
    summary = {
        "generatedAt": generated_at,
        "syncScope": args.sync_scope,
        "registry": {
            "parsed": registry_parsed,
            "records": len(well_records),
            "warnings": len(well_warnings),
        },
        "waterRights": {
            "records": len(attachment_records),
            "warnings": len(attachment_warnings),
        },
        "pumping": pumping_stats,
        "notes": notes,
    }
    write_summary(args.work_dir / INDEX_FILENAMES["summary"], summary)

    well_uploads = [
        INDEX_FILENAMES["well"],
        INDEX_FILENAMES["station"],
        INDEX_FILENAMES["wellWarnings"],
        INDEX_FILENAMES["attachments"],
        INDEX_FILENAMES["sync"],
        INDEX_FILENAMES["summary"],
    ]
    upload_indexes(service, args.well_index_folder_id, well_uploads, args.work_dir, "well", notes)

    pumping_uploads = [
        INDEX_FILENAMES["pumping"],
        INDEX_FILENAMES["pumpingMonth"],
        INDEX_FILENAMES["pumpingWarnings"],
    ]
    upload_indexes(service, args.pumping_index_folder_id, pumping_uploads, args.work_dir, "pumping", notes)

    print(json.dumps(summary, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as exc:
        print(f"groundwater_drive_sync failed: {exc}", file=sys.stderr)
        raise
