"""Read monthly pumping workbooks without changing their contents.

The summary sheet is authoritative; station sheets supply the reporting year.
All validation completes before any index is replaced.
"""
from __future__ import annotations

import hashlib
import json
import math
import re
import unicodedata
from pathlib import Path


def text(value):
    return unicodedata.normalize('NFKC', str(value if value is not None else '')).strip()


def number(value, location):
    if value is None or text(value) in ('', '-', '—', '--', '未填報'):
        return None
    if isinstance(value, bool):
        raise ValueError(f'{location}: boolean is not a pumping amount')
    try:
        value = float(text(value).replace(',', ''))
    except ValueError as exc:
        raise ValueError(f'{location}: invalid pumping amount') from exc
    if not math.isfinite(value) or value < 0:
        raise ValueError(f'{location}: pumping amount must be finite and non-negative')
    return value


def year(value):
    match = re.fullmatch(r'(?:民國)?\s*(\d{2,4})(?:\.0)?\s*(?:年度?|年度別)?', text(value))
    if not match:
        raise ValueError(f'Cannot determine reporting year: {value!r}')
    result = int(match[1])
    if result >= 1912:
        result -= 1911
    if not 1 <= result <= 300:
        raise ValueError('Reporting year out of range')
    return result


def workbook_year(workbook, source_name):
    years = set()
    for sheet in workbook:
        for row in sheet.iter_rows(max_row=8, values_only=True):
            for i, value in enumerate(row[:-1]):
                if text(value).rstrip('-：: ') in ('年度別', '年度', '民國年度'):
                    years.add(year(row[i + 1]))
    if len(years) == 1:
        return years.pop()
    if len(years) > 1:
        raise ValueError('Workbook contains conflicting reporting years; use explicit year column')
    match = re.search(r'(?<!\d)(\d{3})(?:年度?|[-_].*|(?=\d{4}(?!\d)))', source_name)
    if not match:
        raise ValueError('Workbook reporting year is missing')
    return year(match[1])


def month_header(value):
    value = text(value).replace(' ', '')
    chinese = ['一', '二', '三', '四', '五', '六', '七', '八', '九', '十', '十一', '十二']
    for i, label in enumerate(chinese, 1):
        if value in (f'{label}月', f'{i}月', f'{i:02}月'):
            return i
    return None


def parse_workbook(path, source=None):
    import openpyxl
    source = source or {'name': Path(path).name}
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    formulas = openpyxl.load_workbook(path, read_only=True, data_only=False)
    try:
        from openpyxl.utils.cell import coordinate_to_tuple
        value_grids = {}
        formula_grids = {}
        def grid(sheet_name, formula=False):
            cache = formula_grids if formula else value_grids
            if sheet_name not in cache:
                cache[sheet_name] = list((formulas if formula else workbook)[sheet_name].values)
            return cache[sheet_name]
        def is_proven_blank(sheet_name, address, visited=None):
            visited = set() if visited is None else visited
            identity = sheet_name, address.replace('$', '')
            if identity in visited:
                return False
            visited.add(identity)
            r, c = coordinate_to_tuple(identity[1])
            values = grid(sheet_name)
            source_cells = grid(sheet_name, True)
            value = values[r-1][c-1]
            formula = source_cells[r-1][c-1]
            if value not in (None, ''):
                return False
            if formula in (None, ''):
                return True
            if not isinstance(formula, str):
                return False
            # A saved empty-string result looks like an absent formula cache.
            # Prove emptiness only for the source workbook's direct references
            # and IF(input="","",...) guards; never evaluate arbitrary Excel.
            ref = re.fullmatch(r"=(?:'([^']+)'|([^!]+))!([A-Z$]+[0-9]+)", formula)
            if ref:
                return is_proven_blank(ref[1] or ref[2], ref[3], visited)
            guard = re.match(r'=IF\(\s*([A-Z$]+[0-9]+)\s*=\s*""\s*,\s*""\s*,', formula, re.I)
            return bool(guard and is_proven_blank(sheet_name, guard[1], visited))
        # Do not double count station sheets and their summary.
        sheets = [workbook['月實取水量表']] if '月實取水量表' in workbook.sheetnames else list(workbook)
        records = []
        inferred_year = None
        for sheet in sheets:
            rows = grid(sheet.title)
            header = None
            for offset, row in enumerate(rows[:30]):
                labels = {text(v): i for i, v in enumerate(row) if text(v)}
                months = {month_header(v): i for i, v in enumerate(row) if month_header(v)}
                if '水權狀號' in labels and len(months) == 12:
                    header = offset, labels, months
                    break
            if header is None:
                continue
            offset, labels, months = header
            year_col = next((labels[k] for k in ('年度', '年度別', '民國年', 'yearMinguo') if k in labels), None)
            if year_col is None and inferred_year is None:
                inferred_year = workbook_year(workbook, source.get('name', Path(path).name))
            formula_rows = grid(sheet.title, True)
            for row_index, row in enumerate(rows[offset + 1:], offset + 2):
                key = text(row[labels['水權狀號']]).upper().replace(' ', '')
                if not key:
                    # A populated well row must never disappear silently.
                    if any(text(row[labels[k]]) for k in ('名稱', '井名', '水井名稱') if k in labels):
                        raise ValueError(f'{sheet.title}!{row_index}: missing water-right number')
                    continue
                if key in ('水權狀號', '合計', '總計'):
                    continue
                if not re.fullmatch(r'[A-Z]\d{7}', key):
                    raise ValueError(f'{sheet.title}!{row_index}: invalid water-right number {key}')
                def cell(col):
                    value = row[col]
                    formula = formula_rows[row_index - 1][col]
                    address = openpyxl.utils.get_column_letter(col + 1) + str(row_index)
                    if value is None and isinstance(formula, str) and formula.startswith('=') and not is_proven_blank(sheet.title, address):
                        raise ValueError(f'{sheet.title}!{row_index}:{col + 1}: formula has no cached result; recalculate and save Excel')
                    return value
                values = [number(cell(months[m]), f'{sheet.title}!{row_index}:month{m}') for m in range(1, 13)]
                total_col = next((labels[k] for k in labels if k.startswith(('年實取水量', '年度合計', '年合計')) or k == '合計'), None)
                total = number(cell(total_col), f'{sheet.title}!{row_index}:total') if total_col is not None else None
                if total is not None and abs(total - sum(v for v in values if v is not None)) > 0.01:
                    raise ValueError(f'{sheet.title}!{row_index} {key}: source total does not match monthly sum')
                def field(*names):
                    return next((text(row[labels[k]]) for k in names if k in labels), '')
                records.append({
                    'waterRightNo': key, 'yearMinguo': year(cell(year_col)) if year_col is not None else inferred_year,
                    'station': field('站別', '工作站'), 'wellName': field('名稱', '井名', '水井名稱'),
                    'authority': field('主管機關') or {'B': '臺中市政府', 'K': '苗栗縣政府'}.get(key[0], ''),
                    'monthlyM3': values, 'sourceTotalM3': total,
                    'source': {**source, 'sheet': sheet.title, 'row': row_index},
                })
        if not records:
            raise ValueError('No supported monthly pumping table found; existing public data preserved')
        # Also validates duplicate keys and numeric values for the JSON path.
        return merge_records([], records)[0]
    finally:
        workbook.close()
        formulas.close()


def merge_records(previous, incoming):
    """Retain historical years/months and reject conflicting duplicate input."""
    result = {(r['waterRightNo'], r['yearMinguo']): dict(r) for r in previous}
    seen = {}
    warnings = []
    changed = []
    for record in incoming:
        key = text(record.get('waterRightNo')).upper().replace(' ', '')
        if not re.fullmatch(r'[A-Z]\d{7}', key):
            raise ValueError(f'Invalid water-right number: {key}')
        y = year(record.get('yearMinguo'))
        identity = key, y
        months = record.get('monthlyM3')
        if not isinstance(months, list) or len(months) != 12:
            raise ValueError(f'{key}:{y}: expected 12 monthly values')
        months = [number(v, f'{key}:{y}:{i+1}') for i, v in enumerate(months)]
        total = number(record.get('sourceTotalM3'), f'{key}:{y}:total')
        if total is not None and abs(total - sum(v for v in months if v is not None)) > 0.01:
            raise ValueError(f'{key}:{y}: source total does not match monthly sum')
        if identity in seen:
            if seen[identity] != (months, total):
                raise ValueError(f'{key}:{y}: conflicting duplicate records')
            continue
        seen[identity] = months, total
        before = result.get(identity, {})
        old = before.get('monthlyM3', [None] * 12)
        if not isinstance(old, list) or len(old) != 12:
            raise ValueError(f'{key}:{y}: existing record does not have 12 months')
        merged = [a if b is None else b for a, b in zip(old, months)]
        for m, (a, b) in enumerate(zip(old, merged), 1):
            if a != b:
                changed.append({'waterRightNo': key, 'yearMinguo': y, 'month': m, 'before': a, 'after': b})
                if a is not None:
                    warnings.append({'waterRightNo': key, 'yearMinguo': y, 'month': m, 'reason': 'historical_value_changed', 'before': a, 'after': b})
        normalized = {**before, **record, 'waterRightNo': key, 'yearMinguo': y, 'monthlyM3': merged}
        for name in ('station', 'wellName', 'authority'):
            normalized[name] = record.get(name) or before.get(name, '')
        if merged != months:
            normalized['sourceReportedTotalM3'] = total
            normalized['sourceTotalM3'] = sum(v for v in merged if v is not None)
            warnings.append({'waterRightNo': key, 'yearMinguo': y, 'reason': 'retained_previous_nonblank_months'})
        else:
            normalized['sourceTotalM3'] = total
        normalized['yearlyKey'] = f'{key}:{y}'
        payload = {k: normalized.get(k) for k in ('waterRightNo', 'yearMinguo', 'station', 'wellName', 'authority', 'monthlyM3', 'sourceTotalM3')}
        normalized['recordHash'] = hashlib.sha256(json.dumps(payload, ensure_ascii=False, sort_keys=True).encode()).hexdigest()
        result[identity] = normalized
    return [result[k] for k in sorted(result, key=lambda k: (k[0], -k[1]))], warnings, changed
